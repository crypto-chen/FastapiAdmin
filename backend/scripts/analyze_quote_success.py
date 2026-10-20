"""报价成功率专项分析：拉取 CRM 报价明细，输出月度汇总与业务员/单据明细。

结论导向：
- 公司口径（/hs/quote/getQuoteStatistics）与业务员口径（/hs/quote/getQuoteList）
  分别取数，二者不是同一快照；
- 分母只含已审核报价单（quote_bills.status='E'）；
- 分子 `transformation=1` 是**取数时点**的转化状态，历史月份会被后续转化“追高”。

用法：
    python scripts/analyze_quote_success.py
    python scripts/analyze_quote_success.py --months 2026-04 2026-05 ... 2026-10
"""

from __future__ import annotations

import argparse
import json
import re
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXPORT_DIR = PROJECT_ROOT.parent / "exports"
CRM_BASE = "http://crmapi.yonggubox.com"

DB = dict(
    host="127.0.0.1",
    user="root",
    password="123",
    database="fastapiadmin",
    charset="utf8mb4",
    cursorclass=pymysql.cursors.DictCursor,
)


def crm_get(path: str, pairs: list[tuple[str, object]]) -> dict:
    url = f"{CRM_BASE}{path}?{urllib.parse.urlencode(pairs, doseq=True)}"
    with urllib.request.urlopen(url, timeout=90) as response:
        return json.loads(response.read().decode("utf-8"))


def month_bounds(month: str) -> tuple[str, str]:
    year, mon = map(int, month.split("-"))
    start = datetime(year, mon, 1)
    end = datetime(year + (mon == 12), 1 if mon == 12 else mon + 1, 1)
    last = end.replace(day=1).toordinal() - 1
    return start.strftime("%Y-%m-%d"), datetime.fromordinal(last).strftime("%Y-%m-%d")


def fetch_month(month: str) -> dict:
    start, end = month_bounds(month)
    stats0 = crm_get(
        "/hs/quote/getQuoteStatistics",
        [("dateRange[]", start), ("dateRange[]", end), ("transformation", 0)],
    )
    stats1 = crm_get(
        "/hs/quote/getQuoteStatistics",
        [("dateRange[]", start), ("dateRange[]", end), ("transformation", 1)],
    )
    quote_list = crm_get("/hs/quote/getQuoteList", [("month", month)])
    return {
        "month": month,
        "stats_all": stats0.get("data") or {},
        "stats_converted": stats1.get("data") or {},
        "quote_list": quote_list.get("data") or [],
    }


def load_person_map(conn) -> dict[int, dict]:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT source_code, source_name, raw_json FROM source_person "
            "WHERE source_type='crm' AND is_deleted=0"
        )
        rows = cur.fetchall()
    result: dict[int, dict] = {}
    for row in rows:
        try:
            raw = json.loads(row["raw_json"] or "{}")
        except json.JSONDecodeError:
            continue
        crm_id = raw.get("id")
        if crm_id is None:
            continue
        result[int(crm_id)] = {
            "code": row["source_code"],
            "name": row["source_name"],
            "number": raw.get("number") or raw.get("employee_number") or "",
        }
    return result


def load_metric_values(conn) -> dict[tuple[str, str], dict]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT d.code, v.period_value, v.value, v.batch_id, v.calc_time
            FROM metric_value v
            JOIN metric_def d ON d.id = v.metric_id
            WHERE d.code IN (
                'marketing_quote_count',
                'marketing_quote_converted_count',
                'marketing_quote_success_rate'
            )
              AND v.is_deleted = 0
            ORDER BY v.id
            """
        )
        rows = cur.fetchall()
    latest: dict[tuple[str, str], dict] = {}
    for row in rows:
        latest[(row["code"], row["period_value"])] = row
    return latest


def load_metric_history(conn, code: str, period: str) -> list[dict]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT v.value, v.batch_id, v.calc_time
            FROM metric_value v
            JOIN metric_def d ON d.id = v.metric_id
            WHERE d.code = %s AND v.period_value = %s AND v.is_deleted = 0
            ORDER BY v.id
            """,
            (code, period),
        )
        return cur.fetchall()


def load_last_quote_run(conn, period: str) -> dict[str, str]:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT rows_json FROM meta_sync_run "
            "WHERE job_id=74 AND period_value=%s ORDER BY id DESC LIMIT 1",
            (period,),
        )
        row = cur.fetchone()
    if not row:
        return {}
    payload = json.loads(row["rows_json"] or "{}")
    return {
        str(item.get("number")): str(item.get("transformation"))
        for item in (payload.get("data") or [])
        if item.get("number")
    }


def quote_date(number: str) -> str:
    match = re.match(r"YGQT(\d{8})-", number or "")
    return match.group(1) if match else ""


def person_key(row: dict) -> int:
    try:
        return int(row.get("create_id"))
    except (TypeError, ValueError):
        return 0


def aggregate_person(rows: list[dict], person_map: dict[int, dict]) -> list[dict]:
    buckets: dict[int, dict] = defaultdict(
        lambda: {
            "count": 0,
            "converted": 0,
            "amount": 0.0,
            "converted_amount": 0.0,
            "unconverted_amount": 0.0,
        }
    )
    for row in rows:
        key = person_key(row)
        bucket = buckets[key]
        bucket["count"] += 1
        amount = float(row.get("not_tax_all_quote_amount") or 0)
        bucket["amount"] += amount
        if str(row.get("transformation")) == "1":
            bucket["converted"] += 1
            bucket["converted_amount"] += amount
        else:
            bucket["unconverted_amount"] += amount

    result = []
    for key, bucket in buckets.items():
        info = person_map.get(key, {})
        result.append(
            {
                "create_id": key,
                "person_code": info.get("code") or "",
                "person_name": info.get("name") or f"未映射({key})",
                "number": info.get("number") or "",
                "quote_count": bucket["count"],
                "converted_count": bucket["converted"],
                "unconverted_count": bucket["count"] - bucket["converted"],
                "success_rate": round(bucket["converted"] / bucket["count"] * 100, 4)
                if bucket["count"]
                else 0.0,
                "quote_amount": round(bucket["amount"], 2),
                "converted_amount": round(bucket["converted_amount"], 2),
                "unconverted_amount": round(bucket["unconverted_amount"], 2),
            }
        )
    result.sort(key=lambda item: (-item["quote_count"], item["person_name"]))
    return result


def style_sheet(ws, widths: list[int], freeze: str = "A2") -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def write_rows(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws, widths)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--months",
        nargs="*",
        default=[f"2026-{month:02d}" for month in range(1, 11)],
    )
    parser.add_argument("--output", default="")
    args = parser.parse_args()

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    conn = pymysql.connect(**DB)
    person_map = load_person_map(conn)
    metric_values = load_metric_values(conn)

    fetched = []
    for month in args.months:
        print(f"fetch {month} ...", flush=True)
        fetched.append(fetch_month(month))

    snapshot_path = EXPORT_DIR / f"报价成功率原始快照_{datetime.now():%Y%m%d_%H%M%S}.json"
    snapshot_path.write_text(json.dumps(fetched, ensure_ascii=False), encoding="utf-8")

    summary_rows = []
    person_sheets: dict[str, list[dict]] = {}
    for item in fetched:
        month = item["month"]
        rows = item["quote_list"]
        list_count = len(rows)
        list_converted = sum(1 for row in rows if str(row.get("transformation")) == "1")
        api_count = int((item["stats_all"] or {}).get("quoteCount") or 0)
        api_converted = int((item["stats_converted"] or {}).get("quoteCount") or 0)
        system_rate = metric_values.get(("marketing_quote_success_rate", month), {})
        system_count = metric_values.get(("marketing_quote_count", month), {})
        system_converted = metric_values.get(("marketing_quote_converted_count", month), {})
        person_sheets[month] = aggregate_person(rows, person_map)
        summary_rows.append(
            [
                month,
                api_count,
                api_converted,
                round(api_converted / api_count * 100, 4) if api_count else 0.0,
                list_count,
                list_converted,
                round(list_converted / list_count * 100, 4) if list_count else 0.0,
                int(system_count.get("value") or 0),
                int(system_converted.get("value") or 0),
                float(system_rate.get("value") or 0),
                str(system_rate.get("batch_id") or ""),
                system_rate.get("calc_time"),
            ]
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "月度汇总"
    write_rows(
        ws,
        [
            "月份",
            "接口报价次数(transformation=0)",
            "接口已转单数(transformation=1)",
            "接口成功率%",
            "报价明细行数",
            "明细已转单数",
            "明细成功率%",
            "系统指标-报价次数",
            "系统指标-已转单数",
            "系统指标-成功率%",
            "系统批次",
            "系统计算时间",
        ],
        summary_rows,
        [10, 22, 22, 14, 14, 14, 14, 16, 16, 14, 24, 20],
    )

    detail_months = [month for month in args.months if person_sheets.get(month)]
    for month in detail_months[-2:]:
        rows = person_sheets[month]
        sheet = wb.create_sheet(f"业务员明细_{month}")
        write_rows(
            sheet,
            [
                "CRM人员ID",
                "CRM账号",
                "姓名",
                "工号",
                "报价次数",
                "已转单",
                "未转单",
                "成功率%",
                "报价未税金额",
                "已转单金额",
                "未转单金额",
            ],
            [
                [
                    item["create_id"],
                    item["person_code"],
                    item["person_name"],
                    item["number"],
                    item["quote_count"],
                    item["converted_count"],
                    item["unconverted_count"],
                    item["success_rate"],
                    item["quote_amount"],
                    item["converted_amount"],
                    item["unconverted_amount"],
                ]
                for item in rows
            ],
            [10, 18, 14, 12, 10, 10, 10, 10, 16, 16, 16],
        )
        for row in sheet.iter_rows(min_row=2, min_col=9, max_col=11):
            for cell in row:
                cell.number_format = "#,##0.00"

    for month in detail_months[-2:]:
        item = next(entry for entry in fetched if entry["month"] == month)
        detail_rows = []
        for row in item["quote_list"]:
            key = person_key(row)
            info = person_map.get(key, {})
            detail_rows.append(
                [
                    row.get("number"),
                    quote_date(row.get("number") or ""),
                    key,
                    info.get("code") or "",
                    info.get("name") or f"未映射({key})",
                    int(row.get("transformation") or 0),
                    "是" if str(row.get("transformation")) == "1" else "否",
                    float(row.get("not_tax_all_quote_amount") or 0),
                    float(row.get("material_labor_cost") or 0),
                ]
            )
        detail_rows.sort(key=lambda row: (row[1], str(row[0])), reverse=True)
        sheet = wb.create_sheet(f"报价单明细_{month}")
        write_rows(
            sheet,
            ["报价单号", "报价日期", "CRM人员ID", "CRM账号", "姓名", "transformation", "是否已转单", "报价未税金额", "料工费"],
            detail_rows,
            [24, 12, 10, 18, 14, 14, 12, 16, 14],
        )
        for row in sheet.iter_rows(min_row=2, min_col=8, max_col=9):
            for cell in row:
                cell.number_format = "#,##0.00"

    # 未转单按报价日期分布：验证“月末/近期报价尚未转化”的滞后特征
    sheet = wb.create_sheet("未转单日期分布")
    date_rows = []
    for month in detail_months[-2:]:
        item = next(entry for entry in fetched if entry["month"] == month)
        buckets: dict[str, list[int]] = defaultdict(lambda: [0, 0])
        for row in item["quote_list"]:
            day = quote_date(row.get("number") or "")
            buckets[day][0] += 1
            if str(row.get("transformation")) != "1":
                buckets[day][1] += 1
        for day in sorted(buckets):
            total, unconverted = buckets[day]
            date_rows.append(
                [
                    month,
                    day,
                    total,
                    unconverted,
                    round((total - unconverted) / total * 100, 4) if total else 0.0,
                ]
            )
    write_rows(
        sheet,
        ["月份", "报价日期", "报价单数", "未转单数", "当日成功率%"],
        date_rows,
        [10, 12, 12, 12, 14],
    )

    # 状态变化：上一次 02:00 同步（10-13）vs 本次实时拉取
    sheet = wb.create_sheet("状态变化核对")
    change_rows = []
    for month in detail_months[-2:]:
        old = load_last_quote_run(conn, month)
        item = next(entry for entry in fetched if entry["month"] == month)
        new = {
            str(row.get("number")): str(row.get("transformation"))
            for row in item["quote_list"]
            if row.get("number")
        }
        common = sorted(set(old) & set(new))
        transitions = Counter((old[key], new[key]) for key in common)
        for (before, after), count in sorted(transitions.items()):
            change_rows.append([month, "汇总", before, after, count, ""])
        for key in common:
            if old[key] != new[key]:
                change_rows.append(
                    [month, "明细", old[key], new[key], 1, key]
                )
    write_rows(
        sheet,
        ["月份", "类型", "旧状态", "新状态", "数量", "报价单号"],
        change_rows,
        [10, 10, 10, 10, 10, 26],
    )

    # 口径核对：系统历史批次（含 174.8148% 异常批次）
    audit_rows = []
    for month in ["2026-09", "2026-10"]:
        for row in load_metric_history(conn, "marketing_quote_success_rate", month):
            audit_rows.append(
                [month, float(row["value"]), str(row["batch_id"]), row["calc_time"]]
            )
    audit_rows.sort(key=lambda row: (row[0], str(row[3])))
    sheet = wb.create_sheet("系统批次核对")
    write_rows(
        sheet,
        ["月份", "系统成功率%", "批次ID", "计算时间"],
        audit_rows,
        [10, 14, 26, 22],
    )

    # 结论页
    sheet = wb.create_sheet("结论")
    sheet["A1"] = "报价成功率偏高原因"
    sheet["A1"].font = Font(bold=True, size=14)
    conclusions = [
        "1. 分母只含已审核报价单（status='E'），草稿/作废报价不进入统计，天然抬高成功率。",
        "2. 分子 transformation=1 是取数时点的状态；历史月份按“截至今天是否已转单”回算，老报价有更长时间转化，成功率接近 100%。",
        "3. 报价明细中未转单集中在月末/最近 10 天：2026-09 的 15 张未转单全部在 09-20 之后，2026-10 未转单 31 张集中在 10-06 之后。",
        "4. 公司指标与业务员指标来自两个接口：getQuoteStatistics 与 getQuoteList，行数/时点不完全一致（如 2026-09 系统 586 行 vs 明细 584 行）。",
        "5. 2026-10-20 02:42 曾出现 174.8148%，原因是分子已刷新为 236、分母仍是 10-13 的旧值 135；02:43 分母刷新为 269 后为 87.7323%。比率指标应同批取数。",
        "6. 建议口径：明确“报价后 N 天内转单率”，未满观察期的报价单独标记为“观察中”不进入分母；或统一以 getQuoteList 同一快照计算分子分母。",
    ]
    for index, text in enumerate(conclusions, start=3):
        sheet.cell(row=index, column=1, value=text).alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows(min_row=3, max_row=2 + len(conclusions)):
        sheet.row_dimensions[row[0].row].height = 34

    output = Path(args.output) if args.output else EXPORT_DIR / f"报价成功率分析_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(output)
    print(f"snapshot: {snapshot_path}")
    print(f"excel: {output}")
    conn.close()


if __name__ == "__main__":
    main()

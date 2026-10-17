"""导出「标准成本小计 > 对外出货净额」的明细，供业务核对 CRM 订单行成本口径。

输出一个 xlsx，两个 sheet：

1. **产品线×SKU 倒挂**：PRD-16 标准成本小计 > PRD-26 对外出货净额的明细行
2. **订单级明细**：上表涉及的 SKU 对应的订单行，含 CRM 原始成本字段
   （``material_cost`` / ``material_cost_total`` / ``process_cost`` / ``process_cost_total``）
   与订单链路的派生结果（标准单位成本、订单总成本、入库数量），便于判断
   「CRM 填的是单位成本还是整单金额」。

用法：
    python scripts/export_cost_anomaly.py                       # 默认 2026-09
    python scripts/export_cost_anomaly.py --period 2026-09 --out ../exports/成本倒挂明细_2026-09.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import asyncio  # noqa: E402

from sqlalchemy import func, select  # noqa: E402

from app.core.base_model import MappedBase  # noqa: E402
from app.core.database import async_db_session, async_engine  # noqa: E402
from app.modules.metadata.model import MetaSyncJobModel, MetaSyncRunModel, SourceObjectModel  # noqa: E402
from app.modules.metadata.sync import extract_payload_rows  # noqa: E402
from app.modules.order.model import OrderFullChainModel  # noqa: E402
from app.modules.product.model import ProductLineAnalysisModel  # noqa: E402
from app.utils.import_util import ImportUtil  # noqa: E402

ImportUtil.find_models(MappedBase)

ORDER_OBJECTS = ("/hs/order/orderAnalyze?type=1", "/hs/order/orderAnalyze?type=2")


def num(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


async def latest_rows(db, object_code: str) -> list[dict]:
    run_id = (
        await db.execute(
            select(MetaSyncRunModel.id)
            .join(MetaSyncJobModel, MetaSyncJobModel.id == MetaSyncRunModel.job_id)
            .join(SourceObjectModel, SourceObjectModel.id == MetaSyncJobModel.source_object_id)
            .where(SourceObjectModel.code == object_code, MetaSyncRunModel.status == "success")
            .order_by(MetaSyncRunModel.id.desc())
            .limit(1)
        )
    ).scalars().first()
    if run_id is None:
        return []
    payload = (
        await db.execute(select(MetaSyncRunModel.rows_json).where(MetaSyncRunModel.id == run_id))
    ).scalars().first()
    return list(extract_payload_rows(payload) or [])


async def collect(period: str) -> tuple[list[tuple], list[list]]:
    async with async_db_session() as db:
        version = (
            await db.execute(
                select(func.max(ProductLineAnalysisModel.calc_version)).where(
                    ProductLineAnalysisModel.period_value == period
                )
            )
        ).scalar()
        if version is None:
            raise SystemExit(f"期间 {period} 还没有产品线分析结果，请先跑 run_product_build.py")
        rows = (
            await db.execute(
                select(ProductLineAnalysisModel).where(
                    ProductLineAnalysisModel.period_value == period,
                    ProductLineAnalysisModel.calc_version == version,
                    ProductLineAnalysisModel.is_aggregate == 0,
                )
            )
        ).scalars().all()

        anomalies: list[tuple] = []
        sku_set: set[tuple[str, str]] = set()
        for row in rows:
            extra = row.extra or {}
            std_cost = num(extra.get("PRD-16"))
            net = num(extra.get("PRD-26"))
            if std_cost <= 0 or std_cost <= net:
                continue
            unit_std = num(extra.get("PRD-17"))
            avg_price = num(extra.get("PRD-X03"))
            anomalies.append(
                (
                    row.product_line_code,
                    row.sku,
                    int(row.order_count or 0),
                    num(row.ship_qty),
                    net,
                    unit_std,
                    std_cost,
                    round(std_cost - net, 4),
                    round(unit_std / avg_price, 2) if avg_price else None,
                    avg_price or None,
                )
            )
            sku_set.add((row.product_line_code or "", row.sku or ""))
        anomalies.sort(key=lambda item: -(item[7] or 0))

        # 订单级明细：该期间最新版本订单 + CRM 原始成本字段
        order_version = (
            await db.execute(
                select(func.max(OrderFullChainModel.calc_version)).where(
                    OrderFullChainModel.period_value == period
                )
            )
        ).scalar()
        order_rows = (
            (
                await db.execute(
                    select(OrderFullChainModel).where(
                        OrderFullChainModel.period_value == period,
                        OrderFullChainModel.calc_version == order_version,
                    )
                )
            )
            .scalars()
            .all()
        )
        raw: dict[str, dict] = {}
        for object_code in ORDER_OBJECTS:
            for item in await latest_rows(db, object_code):
                number = str(item.get("number") or "").strip()
                if number:
                    raw[number] = item

        detail: list[list] = []
        for row in order_rows:
            extra = row.extra or {}
            # 按 SKU 过滤（产品线口径可能因上卷与订单行 product_series 不同，这里只按 SKU）
            if not any(extra.get("sku") == sku for _line, sku in sku_set):
                continue
            r = raw.get(row.order_number, {})
            qty = num(row.qty)
            intake = num(row.intake_untaxed)
            detail.append(
                [
                    row.order_number,
                    row.extra.get("product_series") if row.extra else "",
                    row.extra.get("sku") if row.extra else "",
                    qty,
                    num(row.ship_qty),
                    max(qty - num(row.ship_qty), 0),
                    round(intake / qty, 4) if qty else None,
                    intake,
                    num(row.net_shipment),
                    num(r.get("material_cost")),
                    num(r.get("material_cost_total")),
                    num(r.get("process_cost")),
                    num(r.get("process_cost_total")),
                    num((row.extra or {}).get("total_cost")),
                    num((row.extra or {}).get("unit_cost")),
                    num((row.extra or {}).get("material_cost_basis")),
                    num((row.extra or {}).get("inbound_qty")),
                    (row.extra or {}).get("cost_source"),
                    row.data_quality_flag,
                ]
            )
        detail.sort(key=lambda item: -(item[7] or 0))
    await async_engine.dispose()
    return anomalies, detail


def pick(value) -> object:
    """把 Decimal / None 等转成可写进 Excel 的标量。"""
    if value is None:
        return None
    if isinstance(value, (int, float, str)):
        return value
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value)


def write_sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            ws.cell(row=r, column=c, value=pick(value))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"


def build_workbook(period: str, anomalies: list[tuple], detail: list[list], out: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "产品线SKU倒挂"
    write_sheet(
        ws1,
        ["产品线", "SKU", "订单数", "出货数量", "对外出货净额", "单位标准成本",
         "标准成本小计", "差额(成本-净额)", "成本/单价倍数", "平均产品单价"],
        [list(item) for item in anomalies],
        [18, 18, 8, 12, 16, 14, 16, 18, 14, 14],
    )

    ws2 = wb.create_sheet("订单级明细")
    write_sheet(
        ws2,
        ["订单号", "产品系列", "SKU", "数量", "出货数量", "未出货数量", "单价(接单未税/数量)",
         "接单未税", "对外出货净额", "CRM material_cost", "CRM material_cost_total",
         "CRM process_cost", "CRM process_cost_total", "订单总成本", "标准单位成本",
         "材料成本(入库量)", "入库数量", "成本来源", "质量标记"],
        detail,
        [18, 22, 14, 10, 12, 12, 20, 14, 16, 16, 20, 16, 20, 16, 14, 18, 12, 14, 30],
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> None:
    parser = argparse.ArgumentParser(description="导出标准成本倒挂明细")
    parser.add_argument("--period", default="2026-09")
    parser.add_argument(
        "--out",
        default="",
        help="输出路径，默认 ../exports/成本倒挂明细_<period>.xlsx",
    )
    args = parser.parse_args()
    out = Path(args.out) if args.out else PROJECT_ROOT.parent / "exports" / f"成本倒挂明细_{args.period}.xlsx"

    anomalies, detail = asyncio.run(collect(args.period))
    build_workbook(args.period, anomalies, detail, out)
    print(f"期间 {args.period}：倒挂 {len(anomalies)} 个「产品线×SKU」、订单明细 {len(detail)} 行")
    print(f"已写出：{out}（{out.stat().st_size / 1024:.1f} KB）")
    for item in anomalies[:5]:
        print(f"  {item[0]} / {item[1]}：成本 {item[6]:,.2f} vs 净额 {item[4]:,.2f}（差 {item[7]:,.2f}，{item[8]}×）")


if __name__ == "__main__":
    main()

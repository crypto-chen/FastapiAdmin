"""导出指标定义 / 指标结果 / 指标结果组织明细到 Excel（供财务确认）。

用法：
    python scripts/export_metric_excel.py --period 2026-07 --period 2026-08 --period 2026-09
    python scripts/export_metric_excel.py --period 2026-09 --output D:/exports/metric.xlsx

生成的 sheet：
- 说明：口径、数据来源、重算时间等交接信息
- 指标目录：全部指标的编码 / 名称 / 口径注释 / 计算方式
- 指标-<期间>：一行一个指标、一列一个组织的结果矩阵（每个期间一个 sheet）
- 明细-<期间>：指标结果的组织 + 核算维度（部门 / 客户 / 店铺）明细（每个期间一个 sheet）
- 重算批次：各指标该期间最新一次计算的批次号与计算时间

口径：每个指标每个期间只取 calc_version 最大的那一批结果；组织列取「不挂核算维度」
的行（即组织合计），明细 sheet 取挂了核算维度的行。金额单位为元，比率为百分数。
"""

import argparse
import asyncio
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.core.base_model import MappedBase  # noqa: E402
from app.core.database import async_db_session, async_engine  # noqa: E402
from app.core.logger import logger  # noqa: E402
from app.modules.masterdata.model import MasterOrgModel  # noqa: E402
from app.modules.metric.model import MetricDefModel, MetricRunModel, MetricValueModel  # noqa: E402
from app.utils.import_util import ImportUtil  # noqa: E402

ImportUtil.find_models(MappedBase)

logger.remove()
logger.add(sys.stderr, level="INFO", format="{time:HH:mm:ss} | {level} | {message}")

TITLE_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FILL = PatternFill("solid", fgColor="4472C4")
SUB_FONT = Font(bold=True, size=11)
AMOUNT_FMT = "#,##0.00"

# 指标口径里的计算方式（measures.kind）中文名
KIND_LABELS = {
    "account_balance_filter": "金蝶科目余额表过滤聚合",
    "sales_outbound_amount": "聚水潭销售出库汇总",
    "api_scalar_amount": "CRM 接口单值 / 比率",
    "production_instock_amount": "金蝶生产入库单结转",
    "ar_aging_provision": "金蝶应收款账龄分析表计提",
    "metric_sum": "指标汇总（组成指标相加）",
    "metric_scale": "指标折算（组成指标 × 系数）",
    "metric_ratio": "指标比率（分子 ÷ 分母）",
    "metric_diff": "指标差额（加项 − 减项）",
}
SOURCE_LABELS = {
    "account_balance_filter": "金蝶云星空",
    "sales_outbound_amount": "聚水潭",
    "api_scalar_amount": "CRM",
    "production_instock_amount": "金蝶云星空",
    "ar_aging_provision": "金蝶云星空",
    "metric_sum": "本系统派生",
    "metric_scale": "本系统派生",
    "metric_ratio": "本系统派生",
    "metric_diff": "本系统派生",
}


def styled(ws, value, *, font=None, fill=None, fmt=None, align=None):
    """只写模式下必须先包成目标 worksheet 的 Cell 才能带样式。

    样式 ID 是按 worksheet 所属 workbook 注册的，所以 ``ws`` 必须是真正要写入的
    sheet，不能用占位 worksheet，否则生成的 xlsx 打不开（样式索引越界）。
    """
    cell = WriteOnlyCell(ws, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    return cell


def text_header(ws, value):
    return styled(
        ws,
        value,
        font=TITLE_FONT,
        fill=TITLE_FILL,
        align=Alignment(horizontal="center", vertical="center"),
    )


def set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


async def load_data(periods):
    """一次性取出指标定义、指标结果（最新版本）、组织字典与计算批次。"""
    async with async_db_session() as db:
        metrics = list(
            (
                await db.execute(
                    select(MetricDefModel)
                    .where(MetricDefModel.is_deleted == False)  # noqa: E712
                    .order_by(MetricDefModel.id)
                )
            )
            .scalars()
            .all()
        )
        orgs = list(
            (
                await db.execute(
                    select(MasterOrgModel.id, MasterOrgModel.code, MasterOrgModel.name).order_by(
                        MasterOrgModel.code
                    )
                )
            ).all()
        )
        latest = (
            select(
                MetricValueModel.metric_id.label("metric_id"),
                MetricValueModel.period_value.label("period_value"),
                func.max(MetricValueModel.calc_version).label("calc_version"),
            )
            .where(MetricValueModel.period_value.in_(periods))
            # 必须按「指标 + 期间」取最新版本：同一指标不同期间的 calc_version 不同，
            # 只按指标分组会把版本号较低的期间整段过滤掉
            .group_by(MetricValueModel.metric_id, MetricValueModel.period_value)
            .subquery()
        )
        rows = (
            await db.execute(
                select(
                    MetricValueModel.metric_id,
                    MetricValueModel.period_value,
                    MetricValueModel.org_id,
                    MetricValueModel.dept_code,
                    MetricValueModel.dept_name,
                    MetricValueModel.value,
                    MetricValueModel.calc_version,
                )
                .join(
                    latest,
                    (MetricValueModel.metric_id == latest.c.metric_id)
                    & (MetricValueModel.period_value == latest.c.period_value)
                    & (MetricValueModel.calc_version == latest.c.calc_version),
                )
                .where(MetricValueModel.period_value.in_(periods))
                .order_by(MetricValueModel.org_id, MetricValueModel.dept_code)
            )
        ).all()
        runs = (
            await db.execute(
                select(
                    MetricRunModel.metric_id,
                    MetricRunModel.period_value,
                    MetricRunModel.version,
                    MetricRunModel.batch_id,
                    MetricRunModel.status,
                    MetricRunModel.finished_at,
                ).where(MetricRunModel.period_value.in_(periods))
            )
        ).all()
    return metrics, orgs, rows, runs


def build_sheets(workbook, periods, metrics, orgs, rows, runs, generated_at):
    metric_by_id = {metric.id: metric for metric in metrics}
    org_by_id = {row[0]: (row[1], row[2]) for row in orgs}
    # 结果行已按「最新 calc_version」过滤；这里拆成组织合计与核算维度明细两组
    totals = {}
    details = {period: [] for period in periods}
    for metric_id, period, org_id, dept_code, dept_name, value, calc_version in rows:
        if dept_code:
            details[period].append(
                (metric_id, org_id, dept_code, dept_name, float(value or 0), calc_version)
            )
        else:
            bucket = totals.setdefault((metric_id, period), {"orgs": {}, "company": 0.0})
            if org_id is None:
                bucket["company"] += float(value or 0)
            else:
                bucket["orgs"][org_id] = bucket["orgs"].get(org_id, 0.0) + float(value or 0)

    sheet_names = []

    # 工时指标（CRM 企微打卡）在接口无数据的月份会返回 0，这里自动提示，避免财务误判
    hour_metrics = [
        metric for metric in metrics if metric.code in ("marketing_regular_work_hours", "marketing_overtime_hours")
    ]
    zero_hour_periods = [
        period
        for period in periods
        if hour_metrics
        and all(
            round(sum((totals.get((metric.id, period)) or {}).get("orgs", {}).values())
                  + (totals.get((metric.id, period)) or {}).get("company", 0.0), 4)
            == 0
            for metric in hour_metrics
        )
    ]

    # ---------- 说明 ----------
    ws = workbook.create_sheet("说明")
    set_widths(ws, [22, 110])
    lines = [
        ("生成时间", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("数据期间", "、".join(periods)),
        ("指标数量", f"{len(metrics)} 个（含派生指标）"),
        ("组织范围", "、".join(f"{code} {name}" for _id, code, name in orgs)),
        (
            "取数口径",
            "每个指标每个期间只取最新一次计算（calc_version 最大）的结果；「指标-<期间>」sheet "
            "的组织列为该组织的合计值（不挂核算维度的行），「明细-<期间>」sheet 为挂了核算维度"
            "（部门 / 客户 / 店铺）的行。",
        ),
        (
            "数据来源",
            "金蝶云星空（科目余额表 / 生产入库单 / 应收款账龄分析表）、聚水潭（销售出库单）、"
            "CRM（订单、出货、下推接口）",
        ),
        ("金额单位", "元；正负号沿用系统口径（如退货 / 退款、应收余额贷方为负数）"),
        ("比率类指标", "报价成功率、报价毛利率、下推率、制造成本率等为百分数（已 ×100）"),
        (
            "汇总行含义",
            "「公司整体」列 = 不按组织拆分的公司级指标（CRM / 聚水潭系列、备货订单入库、总收益、"
            "应收周转天数等）；「合计」列 = 各组织 + 公司整体之和",
        ),
        (
            "空值说明",
            "「指标-<期间>」里空单元格表示该指标当月在该组织无结果；「明细-<期间>」里组织列为空"
            "表示该指标未按组织拆分（公司整体口径），核算维度列为订单号 / 客户 / 部门 / 店铺等",
        ),
        (
            "待确认事项",
            "账龄计提（坏账准备）取《应收款账龄分析表》按往来单位净值口径，与「营销中心应收余额」"
            "（科目余额表 1122）口径不同，需财务确认计提基数",
        ),
    ]
    if zero_hour_periods and len(zero_hour_periods) < len(periods):
        lines.append(
            (
                "已知数据缺口",
                f"营销中心正常工作时间 / 加班时间（CRM 企微打卡接口 /hs/getWorkTime）在 "
                f"{'、'.join(zero_hour_periods)} 返回 0（接口该期间未返回有效打卡数据），不是真实工时；"
                f"其余月份为接口实际值",
            )
        )
    for label, value in lines:
        ws.append([styled(ws, label, font=SUB_FONT), styled(ws, value)])
    sheet_names.append("说明")

    # ---------- 指标目录 ----------
    ws = workbook.create_sheet("指标目录")
    headers = [
        "指标编码",
        "指标名称",
        "Excel科目编码",
        "期间类型",
        "计算方式",
        "数据来源",
        "口径说明",
        "公式",
        "敏感级别",
        "状态",
        "版本",
    ]
    ws.append([text_header(ws, item) for item in headers])
    set_widths(ws, [34, 34, 14, 12, 24, 14, 90, 40, 10, 8, 8])
    for metric in metrics:
        kind = str((metric.measures or {}).get("kind") or "")
        ws.append(
            [
                metric.code,
                metric.name,
                metric.excel_code or "-",
                metric.period_type,
                KIND_LABELS.get(kind, kind or "-"),
                SOURCE_LABELS.get(kind, "-"),
                metric.description or "-",
                metric.formula or "-",
                metric.sensitivity,
                "启用" if metric.status == 0 else "停用",
                metric.version,
            ]
        )
    sheet_names.append("指标目录")

    # ---------- 每月指标矩阵 ----------
    for period in periods:
        ws = workbook.create_sheet(f"指标-{period}")
        org_columns = [(row[0], row[1], row[2]) for row in orgs]
        headers = (
            ["指标编码", "指标名称", "口径说明"]
            + [f"{code} {name}" for _id, code, name in org_columns]
            + ["公司整体", "合计"]
        )
        ws.append([text_header(ws, item) for item in headers])
        set_widths(ws, [34, 34, 60] + [20] * len(org_columns) + [18, 18])
        ws.freeze_panes = "D2"
        for metric in metrics:
            bucket = totals.get((metric.id, period))
            cells = [metric.code, metric.name, metric.description or "-"]
            for org_id, _code, _name in org_columns:
                cells.append(bucket["orgs"].get(org_id) if bucket else None)
            cells.append(bucket["company"] if bucket else None)
            cells.append(
                round(sum(bucket["orgs"].values()) + bucket["company"], 4) if bucket else None
            )
            ws.append(
                [
                    styled(ws, item, fmt=AMOUNT_FMT) if isinstance(item, float) else item
                    for item in cells
                ]
            )
        sheet_names.append(ws.title)

    # ---------- 每月明细 ----------
    for period in periods:
        ws = workbook.create_sheet(f"明细-{period}")
        headers = [
            "指标编码",
            "指标名称",
            "组织编码",
            "组织名称",
            "核算维度编码",
            "核算维度名称",
            "指标值",
            "计算版本",
        ]
        ws.append([text_header(ws, item) for item in headers])
        set_widths(ws, [34, 34, 12, 34, 22, 34, 18, 10])
        ws.freeze_panes = "A2"
        for metric_id, org_id, dept_code, dept_name, value, calc_version in details[period]:
            metric = metric_by_id.get(metric_id)
            org_code, org_name = org_by_id.get(org_id, ("", ""))
            ws.append(
                [
                    metric.code if metric else str(metric_id),
                    metric.name if metric else "",
                    org_code,
                    org_name,
                    dept_code,
                    dept_name or "",
                    styled(ws, value, fmt=AMOUNT_FMT),
                    calc_version,
                ]
            )
        sheet_names.append(ws.title)

    # ---------- 重算批次 ----------
    ws = workbook.create_sheet("重算批次")
    headers = ["指标编码", "指标名称", "期间", "指标版本", "计算批次号", "状态", "完成时间"]
    ws.append([text_header(ws, item) for item in headers])
    set_widths(ws, [34, 34, 12, 10, 26, 12, 22])
    latest_run = {}
    for metric_id, period, version, batch_id, status, finished_at in runs:
        key = (metric_id, period)
        current = latest_run.get(key)
        if current is None or (finished_at or datetime.min) > (current[3] or datetime.min):
            latest_run[key] = (version, batch_id, status, finished_at)
    for metric in metrics:
        for period in periods:
            item = latest_run.get((metric.id, period))
            ws.append(
                [
                    metric.code,
                    metric.name,
                    period,
                    item[0] if item else None,
                    item[1] if item else "",
                    item[2] if item else "无计算记录",
                    item[3].strftime("%Y-%m-%d %H:%M:%S") if item and item[3] else "",
                ]
            )
    sheet_names.append("重算批次")

    return sheet_names


async def main():
    parser = argparse.ArgumentParser(description="导出指标结果到 Excel")
    parser.add_argument("--period", action="append", required=True, help="期间值，如 2026-09，可重复")
    parser.add_argument("--output", default=None, help="输出文件路径（缺省写到项目 exports 目录）")
    args = parser.parse_args()

    periods = list(dict.fromkeys(args.period))
    generated_at = datetime.now()
    metrics, orgs, rows, runs = await load_data(periods)
    print(f"指标 {len(metrics)} 个，结果行 {len(rows)} 行，组织 {len(orgs)} 个", flush=True)

    workbook = Workbook(write_only=True)
    sheet_names = build_sheets(workbook, periods, metrics, orgs, rows, runs, generated_at)

    if args.output:
        output = Path(args.output)
    else:
        stamp = generated_at.strftime("%Y%m%d_%H%M")
        span = f"{periods[0]}~{periods[-1]}" if len(periods) > 1 else periods[0]
        output = PROJECT_ROOT.parent / "exports" / f"核算指标_{span}_{stamp}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"已生成: {output}")
    print(f"sheet: {', '.join(sheet_names)}")
    await async_engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
